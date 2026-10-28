"""
template_d.py — Chuẩn hóa file Mẫu 03 về ĐÚNG khuôn "mẫu D" (keyword mẫu 3 cột FY): cột A..FY,
sắp xếp theo MÃ FIELD (keyword ở dòng mã), bất kể file upload đang theo mẫu cũ (cột C) hay lệch cột.

Quy tắc:
  * Cột nào có keyword thuộc mẫu D -> đặt đúng vị trí mẫu D (kèm dữ liệu, định dạng, công thức).
  * Keyword mẫu cũ 'noi_cong_tac' (cột S) -> đổi thành 'noi_cong_tac_id' (nếu file chưa có sẵn
    'noi_cong_tac_id').
  * File thiếu keyword nào của mẫu D -> tự thêm cột trống đúng vị trí (nhãn lấy ở NEW_COLUMN_LABELS).
  * Cột có keyword KHÔNG thuộc mẫu D (vd 'nghenghiep_id', 'noi_cong_tac' ở GE/GF, các cột ngoài FY...)
    -> bỏ. Cột T (mẫu D không có keyword) giữ nguyên nếu file nguồn cũng có 1 cột không keyword ở T.
"""
from copy import copy
from io import BytesIO
import re

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange

from validators import SHEET_MAIN, detect_layout_rows, clean_ws

# Index 0 = cột A ... index 180 = cột FY. None = cột không có keyword (A = STT, T = cột phụ).
TEMPLATE_D_COLUMNS = [
    None,
    'ma_phieu',
    'ngay_kham',
    'doi_tuong_kham',
    'dia_diem_kham',
    'dinh_danh_ca_nhan',
    'ho_ten',
    'ngay_sinh',
    'gioi_tinh',
    'dan_toc_id',
    'nhom_mau_id',
    'yeu_to_nhom_mau_id',
    'the_bhyt',
    'sdt',
    'dia_chi_hien_tai',
    'city_id',
    'ward_id',
    'nghenghiep_code',
    'noi_cong_tac_id',
    None,
    'noi_cong_tac_xa_phuong',
    'ly_do_kham',
    'tinhtrangthucnghiem_xetnghiem',
    'hinh_thuc_chi_tra_khamsk',
    'hinh_thuc_chi_tra_khamsk_chi_tiet',
    'nguonkhac_ghiro',
    'giadinh_macbenh',
    'giadinh_danhsachbenh_icd',
    'benh_5nam',
    'benh_than_kinh',
    'benh_mat',
    'benh_tai',
    'benh_tim',
    'pt_tim_mach',
    'tang_ha',
    'kho_tho',
    'benh_phoi',
    'benh_than',
    'nghien_ruou_bia',
    'dai_thao_duong',
    'benh_tam_than',
    'mat_y_thuc',
    'ngat_chong_mat',
    'benh_tieu_hoa',
    'roi_loan_giac_ngu',
    'tai_bien_mach_mau_nao',
    'cot_song',
    'su_dung_ruou_bia',
    'su_dung_ma_tuy',
    'benh_khac',
    'dieu_tri_benh_co_khong',
    'dieu_tri_benh_liet_ke',
    'thai_san_co_khong',
    'thai_san_liet_ke',
    'chieucao',
    'cannang',
    'nhiptho',
    'mach',
    'huyetaptamthu',
    'huyetaptamtruong',
    'phanloai',
    'noikhoa_chuaphathienbatthuong',
    'noikhoa_chandoansobo_icd',
    'noikhoa_chandoanxacdinh_icd',
    'noikhoa_phanloai',
    'hohap_chuaphathienbatthuong',
    'hohap_chandoansobo_icd',
    'hohap_chandoanxacdinh_icd',
    'hohap_phanloai',
    'tieuhoa_chuaphathienbatthuong',
    'tieuhoa_chandoansobo_icd',
    'tieuhoa_chandoanxacdinh_icd',
    'tieuhoa_phanloai',
    'thantietnieu_chuaphathienbatthuong',
    'thantietnieu_chandoansobo_icd',
    'thantietnieu_chandoanxacdinh_icd',
    'thantietnieu_phanloai',
    'noitiet_chuaphathienbatthuong',
    'noitiet_chandoansobo_icd',
    'noitiet_chandoanxacdinh_icd',
    'noitiet_phanloai',
    'coxuongkhop_chuaphathienbatthuong',
    'coxuongkhop_chandoansobo_icd',
    'coxuongkhop_chandoanxacdinh_icd',
    'coxuongkhop_phanloai',
    'thankinh_chuaphathienbatthuong',
    'thankinh_chuandoansobo_icd',
    'thankinh_chandoanxacdinh_icd',
    'thankinh_phanloai',
    'tamthan_chuaphathienbatthuong',
    'tamthan_chandoansobo_icd',
    'tamthan_chandoanxacdinh_icd',
    'tamthan_phanloai',
    'ngoaikhoa_chuaphathienbatthuong',
    'ngoaikhoa_chandoansobo_icd',
    'ngoaikhoa_chandoanxacdinh_icd',
    'ngoaikhoa_phanloai',
    'dalieu_chuaphathienbatthuong',
    'dalieu_chandoansobo_icd',
    'dalieu_chandoanxacdinh_icd',
    'dalieu_phanloai',
    'sankhoa_tuchoikham',
    'sankhoa_chuaphathienbatthuong',
    'sankhoa_chandoansobo_icd',
    'sankhoa_chandoanxacdinh_icd',
    'sankhoa_phanloai',
    'phukhoa_tuchoikham',
    'phukhoa_chuaphathienbatthuong',
    'phukhoa_chandoansobo_icd',
    'phukhoa_chandoanxacdinh_icd',
    'phukhoa_phanloai',
    'mat_chuaphathienbatthuong',
    'mat_chandoansobo_icd',
    'mat_chandoanxacdinh_icd',
    'mat_phanloai',
    'tmh_chuaphathienbatthuong',
    'tmh_chandoansobo_icd',
    'tmh_chandoanxacdinh_icd',
    'tmh_phanloai',
    'rhm_chuaphathienbatthuong',
    'rhm_chandoansobo_icd',
    'rhm_chandoanxacdinh_icd',
    'rhm_phanloai',
    'mat_khongkinh_mp',
    'mat_khongkinh_mt',
    'mat_kinhlo_mp',
    'mat_kinhlo_mt',
    'mat_cokinh_mp',
    'mat_cokinh_mt',
    'mat_docau_mp',
    'mat_dotru_mp',
    'mat_truc_mp',
    'mat_docau_mt',
    'mat_dotru_mt',
    'mat_truc_mt',
    'tai_noithuong_tt',
    'tai_noithuong_tp',
    'tai_noitham_tt',
    'tai_noitham_tp',
    'loai_kham',
    'kskdk_xnm_slhc',
    'kskdk_xnm_huyetsacto',
    'kskdk_xnm_hematocrit',
    'kskdk_xnm_mcv',
    'kskdk_xnm_mch',
    'kskdk_xnm_mchc',
    'kskdk_xnm_rdw',
    'kskdk_xnm_slbc',
    'kskdk_xnm_slbc_trungtinh',
    'kskdk_xnm_slbc_lympho',
    'kskdk_xnm_slbc_donnhan',
    'kskdk_xnm_slbc_aitoan',
    'kskdk_xnm_slbc_aikiem',
    'kskdk_xnm_sltc',
    'kskdk_shm_duongmau',
    'kskdk_shm_duongmau_lucdoi',
    'kskdk_shm_ure',
    'kskdk_shm_creatinin',
    'kskdk_shm_asat_got',
    'kskdk_shm_alat_gpt',
    'kskdk_xnnt_titrong',
    'kskdk_xnnt_ph',
    'kskdk_xnnt_bachcau',
    'kskdk_xnnt_hongcau',
    'kskdk_xnnt_nitrit',
    'kskdk_xnnt_protein',
    'kskdk_xnnt_glucose',
    'kskdk_xnnt_cetonic',
    'kskdk_xnnt_bilirubin',
    'kskdk_xnnt_urobilinogen',
    'kskdk_xnnt_khac',
    'kskdk_chuan_doan_hinh_anh',
    'can_lam_sang_khac',
    'can_lam_sang_khac_chi_tiet',
    'xet_nghiem_te_bao_co_tu_cung',
    'xet_nghiem_hpv',
    'xquang_nhu',
    'sieu_am_2_tuyen_vu',
    'danh_muc_de_nghi',
    'denghi_chuyentuyen_ghiro',
    'de_nghi'
]
N_COLS = len(TEMPLATE_D_COLUMNS)  # 181 = A..FY
assert N_COLS == column_index_from_string("FY")
TEMPLATE_D_SET = {k for k in TEMPLATE_D_COLUMNS if k}
SPACER_COL = 20  # cột T

# Keyword mẫu cũ -> keyword mẫu D.
RENAMES = {"noi_cong_tac": "noi_cong_tac_id"}

# Cột mẫu D mà mẫu 3 cũ không có -> khi ghép/chuẩn hóa file cũ thì để trống.
NEW_TEMPLATE_ONLY_FIELDS = ("tinhtrangthucnghiem_xetnghiem", "denghi_chuyentuyen_ghiro")


def apply_renames(code_to_col):
    """Gộp keyword mẫu cũ vào keyword mẫu D trên dict {keyword: số cột}, sửa trực tiếp và trả về
    danh sách [(keyword cũ, keyword mới)] đã đổi.

    Mẫu 3 cũ có SẴN cả 'noi_cong_tac' (cột S — ô nhập thật) lẫn 'noi_cong_tac_id' (cột GD — cột phụ
    công thức VLOOKUP). Mẫu mới thì ngược lại: 'noi_cong_tac_id' ở cột S là ô nhập thật còn
    'noi_cong_tac' ở GF là cột phụ. Vì vậy khi file có CẢ HAI keyword thì lấy cột nằm BÊN TRÁI hơn
    (ô nhập thật) làm 'noi_cong_tac_id'; chỉ có 'noi_cong_tac' thì đổi tên."""
    renamed = []
    for old, new in RENAMES.items():
        if old not in code_to_col:
            continue
        if new not in code_to_col or code_to_col[old] < code_to_col[new]:
            code_to_col[new] = code_to_col.pop(old)
            renamed.append((old, new))
    return renamed

# Nhãn dòng 2 (nhãn) / dòng 3 (nhãn phụ) cho cột phải tự thêm. Sửa lại cho đúng nhãn thật của mẫu Medinet.
# keyword: (nhãn, nhãn phụ)
NEW_COLUMN_LABELS = {
    "ma_phieu": ("Mã phiếu", None),
    "tai_noithuong_tt": ("Kết quả khám thính lực", "Tai trái (Nói thường)"),
    "tai_noithuong_tp": ("Kết quả khám thính lực", "Tai phải (Nói Thường)"),
    "tai_noitham_tt": ("Kết quả khám thính lực", "Tai trái (Nói thầm)"),
    "tai_noitham_tp": ("Kết quả khám thính lực", "Tai phải (Nói thầm)"),
    "kskdk_shm_duongmau_lucdoi": ("Đường máu lúc đói (mmol/L)", None),
    "noi_cong_tac_id": ("Nơi công tác", None),
    "tinhtrangthucnghiem_xetnghiem": ("Tình trạng thực nghiệm xét nghiệm", None),
    "denghi_chuyentuyen_ghiro": ("Đề nghị chuyển tuyến (ghi rõ)", None),
}


def _remap_ref(text, mapping, context_sheet):
    """Đổi chữ cái cột trong địa chỉ của sheet chính theo mapping {cột cũ: cột mới}."""
    if "!" in text:
        prefix, addr = text.rsplit("!", 1)
        if prefix.strip("'").replace("''", "'") != SHEET_MAIN:
            return text
        lead = prefix + "!"
    else:
        if context_sheet != SHEET_MAIN:
            return text
        lead, addr = "", text
    if not re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?|\$?[A-Z]{1,3}:\$?[A-Z]{1,3}", addr):
        return text

    def repl(m):
        col = column_index_from_string(m.group(2))
        return m.group(1) + get_column_letter(mapping.get(col, col))
    return lead + re.sub(r"(\$?)([A-Z]{1,3})(?=\$?\d|:|$)", repl, addr)


def _remap_formula(formula, mapping, context_sheet):
    if not formula or not isinstance(formula, str):
        return formula
    has_equal = formula.startswith("=")
    tokens = Tokenizer(formula if has_equal else "=" + formula).items
    out = "".join(_remap_ref(t.value, mapping, context_sheet)
                  if t.type == "OPERAND" and t.subtype == "RANGE" else t.value for t in tokens)
    return ("=" if has_equal else "") + out


def _remap_sqref(sqref, mapping):
    """Đổi vùng áp dụng (data validation / conditional formatting) theo mapping; bỏ cột đã bị xoá."""
    parts = []
    for rng in str(sqref).split():
        cr = CellRange(rng)
        for c in range(cr.min_col, cr.max_col + 1):
            if c in mapping:
                parts.append(f"{get_column_letter(mapping[c])}{cr.min_row}:{get_column_letter(mapping[c])}{cr.max_row}")
    return " ".join(parts)


def _template_layout(ws, code_row):
    """Trả về (header_matches_template: bool, dst_to_src: {cột đích: cột nguồn}, renamed, dropped, added)."""
    src_codes = {}
    for c in range(1, ws.max_column + 1):
        code = clean_ws(ws.cell(row=code_row, column=c).value)
        if code and code not in src_codes:
            src_codes[code] = c
    renamed = apply_renames(src_codes)
    dst_to_src, added = {}, []
    for dst, code in enumerate(TEMPLATE_D_COLUMNS, start=1):
        if code and code in src_codes:
            dst_to_src[dst] = src_codes[code]
        elif code:
            added.append(code)
        elif dst == 1:  # STT: lấy cột 1 nguồn nếu cột đó không mang keyword nào của mẫu
            first = clean_ws(ws.cell(row=code_row, column=1).value)
            if not first or first not in TEMPLATE_D_SET:
                dst_to_src[1] = 1
        elif dst == SPACER_COL:  # cột T: giữ nếu nguồn cũng có cột không keyword ở T
            if ws.max_column >= SPACER_COL and not clean_ws(ws.cell(row=code_row, column=SPACER_COL).value) \
                    and SPACER_COL not in dst_to_src.values():
                dst_to_src[SPACER_COL] = SPACER_COL
    used = set(dst_to_src.values())
    dropped = []
    for c in range(1, ws.max_column + 1):
        code = clean_ws(ws.cell(row=code_row, column=c).value)
        if c not in used and code:
            dropped.append((get_column_letter(c), code))
    already = (not added and not dropped and not renamed
               and all(dst_to_src.get(d, d) == d for d in dst_to_src) and ws.max_column <= N_COLS)
    return already, dst_to_src, renamed, dropped, added


def normalize_to_template_d(file_bytes):
    """Đưa sheet ThongTinHanhChinh về đúng khuôn mẫu D (A..FY) theo keyword.
    Trả về (bytes, info) — info = {changed, renamed, dropped, added}; không cần đổi thì trả nguyên bytes."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes))
    if SHEET_MAIN not in wb.sheetnames:
        return file_bytes, {"changed": False, "renamed": [], "dropped": [], "added": []}
    ws = wb[SHEET_MAIN]
    if ws.tables:
        raise ValueError("Sheet dữ liệu có Excel Table — không tự sắp lại cột được; hãy dùng mẫu Medinet thường.")
    _label_row, code_row, _start, _note = detect_layout_rows(ws)
    already, dst_to_src, renamed, dropped, added = _template_layout(ws, code_row)
    info = {"changed": not already, "renamed": renamed, "dropped": dropped, "added": added}
    if already:
        return file_bytes, info

    src_to_dst = {s: d for d, s in dst_to_src.items()}
    max_row, max_col = ws.max_row, ws.max_column

    # --- snapshot (gỡ merge trước để giữ giá trị ô trên cùng bên trái) ---
    merged = [CellRange(str(r)) for r in ws.merged_cells.ranges]
    for r in merged:
        ws.unmerge_cells(str(r))
    snap = {}
    for r in range(1, max_row + 1):
        for c in src_to_dst:
            cell = ws.cell(row=r, column=c)
            snap[(r, c)] = (cell.value, cell.data_type, copy(cell._style), cell.comment, cell.hyperlink)
    widths = {}
    for key, dim in list(ws.column_dimensions.items()):
        lo = dim.min or column_index_from_string(key)
        hi = dim.max or lo
        for c in range(lo, hi + 1):
            if c in src_to_dst:
                widths[c] = (dim.width, dim.hidden, dim.customWidth)
    dvs = list(ws.data_validations.dataValidation)
    cfs = [(cf.sqref, list(cf.rules)) for cf in ws.conditional_formatting]
    af_ref = ws.auto_filter.ref

    # --- xoá toàn bộ ô cũ rồi ghi lại theo vị trí mới ---
    for key in list(ws._cells):
        del ws._cells[key]
    ws.column_dimensions.clear()

    for dst in range(1, N_COLS + 1):
        src = dst_to_src.get(dst)
        if src:
            for r in range(1, max_row + 1):
                value, dtype, style, comment, link = snap[(r, src)]
                cell = ws.cell(row=r, column=dst)
                if dtype == "f":
                    value = _remap_formula(value, src_to_dst, SHEET_MAIN)
                cell.value = value
                if dtype in ("f", "s") and value is not None:
                    cell.data_type = dtype
                cell._style = style
                if comment:
                    cell.comment = copy(comment)
                if link:
                    cell._hyperlink = copy(link)
            if TEMPLATE_D_COLUMNS[dst - 1]:  # đồng bộ keyword mẫu D (vd noi_cong_tac -> noi_cong_tac_id)
                ws.cell(row=code_row, column=dst).value = TEMPLATE_D_COLUMNS[dst - 1]
            w = widths.get(src)
            if w and (w[0] or w[1]):
                d = ws.column_dimensions[get_column_letter(dst)]
                if w[0]:
                    d.width = w[0]
                d.hidden = bool(w[1])
        else:
            code = TEMPLATE_D_COLUMNS[dst - 1]
            left = dst - 1
            for r in range(1, max_row + 1):
                cell = ws.cell(row=r, column=dst)
                ref = ws.cell(row=r, column=left) if left >= 1 else None
                if ref is not None:
                    cell._style = copy(ref._style)
            if code:
                ws.cell(row=code_row, column=dst).value = code
                label, sub = NEW_COLUMN_LABELS.get(code, (code, None))
                if code_row >= 3:
                    ws.cell(row=code_row - 2, column=dst).value = label
                    ws.cell(row=code_row - 1, column=dst).value = sub
                    if sub is None:
                        ws.merge_cells(start_row=code_row - 2, end_row=code_row - 1, start_column=dst, end_column=dst)

    # --- gộp ô: chỉ giữ lại khi vị trí mới còn liền kề đúng thứ tự ---
    for r in merged:
        cols = list(range(r.min_col, r.max_col + 1))
        if all(c in src_to_dst for c in cols):
            new = [src_to_dst[c] for c in cols]
            if new == list(range(new[0], new[0] + len(new))):
                ref = f"{get_column_letter(new[0])}{r.min_row}:{get_column_letter(new[-1])}{r.max_row}"
                if len(new) > 1 or r.max_row > r.min_row:
                    ws.merge_cells(ref)

    # --- data validation / conditional formatting / autofilter ---
    for dv in dvs:
        dv.sqref = MultiCellRange(_remap_sqref(dv.sqref, src_to_dst))
        dv.formula1 = _remap_formula(dv.formula1, src_to_dst, SHEET_MAIN)
        dv.formula2 = _remap_formula(dv.formula2, src_to_dst, SHEET_MAIN)
    ws.data_validations.dataValidation = [dv for dv in dvs if str(dv.sqref)]
    ws.conditional_formatting = ConditionalFormattingList()
    for sqref, rules in cfs:
        new_ref = _remap_sqref(sqref, src_to_dst)
        if new_ref:
            for rule in rules:
                ws.conditional_formatting.add(new_ref, rule)
    if af_ref:
        cr = CellRange(af_ref)
        ws.auto_filter.ref = f"A{cr.min_row}:{get_column_letter(N_COLS)}{cr.max_row}"

    out = BytesIO()
    wb.save(out)
    return out.getvalue(), info
