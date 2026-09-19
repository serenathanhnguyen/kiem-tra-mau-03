"""
merger.py — Ghép nhiều file Excel Mẫu 03 (mỗi file 1 khóm/tổ/phường) thành 1 file duy nhất,
giữ nguyên cấu trúc mẫu Medinet, để chạy 1 lần bằng script Tampermonkey thay vì chạy tay từng
file nhỏ.

Cố tình TÁI DÙNG logic đọc cấu trúc file (dò dòng nhãn/dòng mã field/dòng dữ liệu, dọn khoảng
trắng ẩn, tìm dòng dữ liệu cuối...) từ validators.py — để không có 2 nơi định nghĩa khác nhau
"thế nào là 1 file Mẫu 03 hợp lệ". Đọc theo MÃ FIELD (không theo số cột), nên vẫn ghép đúng dù
1 file nguồn có cột lệch vị trí đôi chút so với file nền.
"""
import os
from io import BytesIO

import openpyxl
from openpyxl.comments import Comment

from validators import (
    SHEET_MAIN,
    CODE_ROW_ANCHOR_FIELDS,
    detect_layout_rows,
    _build_code_col_map,
    clean_ws,
    is_blank,
    find_last_data_row,
    FILL_WARN,
)

# Thiếu 1 trong các mã "mỏ neo" này thì coi file KHÔNG đúng cấu trúc Mẫu 03 (ví dụ lỡ tải nhầm
# file Mẫu 01/02/04) — dừng hẳn file đó, không đoán mò để tránh ghép sai dữ liệu.
REQUIRED_ANCHOR_FIELDS = CODE_ROW_ANCHOR_FIELDS  # {"ho_ten", "dinh_danh_ca_nhan", "ngay_kham", "gioi_tinh"}


def default_unit_label(filename):
    """Đoán tên đơn vị mặc định từ tên file, để người dùng sửa lại trên giao diện nếu cần."""
    return clean_ws(os.path.splitext(filename)[0])


def read_source_file(file_bytes, filename, unit_label):
    """Đọc 1 file nguồn, trả về dict mô tả cấu trúc + toàn bộ dòng bệnh nhân (map theo MÃ FIELD).
    Raise ValueError với thông báo rõ ràng nếu file không nhận diện được là đúng cấu trúc Mẫu 03."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    if SHEET_MAIN not in wb.sheetnames:
        raise ValueError(f"Không có sheet '{SHEET_MAIN}' trong file — có thể đây không phải file Mẫu 03.")
    ws = wb[SHEET_MAIN]

    label_row, code_row, data_start_row, layout_note = detect_layout_rows(ws)
    col_map = _build_code_col_map(ws, code_row)

    missing_anchor = REQUIRED_ANCHOR_FIELDS - set(col_map)
    if missing_anchor:
        raise ValueError(
            f"Không nhận diện được các mã field bắt buộc {sorted(missing_anchor)} trong dòng mã "
            f"field (dòng {code_row}) — có thể file này không đúng cấu trúc Mẫu 03, hoặc là mẫu khác."
        )

    ho_ten_col = col_map.get("ho_ten")
    cccd_col = col_map.get("dinh_danh_ca_nhan")
    key_cols = [c for c in (ho_ten_col, cccd_col) if c]
    last_row = find_last_data_row(ws, key_cols, data_start_row)

    rows = []
    for r in range(data_start_row, last_row + 1):
        if all(is_blank(ws.cell(row=r, column=c).value) for c in key_cols):
            continue  # dòng trống xen giữa — bỏ qua, không tính là 1 bệnh nhân
        values = {code: ws.cell(row=r, column=c).value for code, c in col_map.items()}
        rows.append({
            "ho_ten": clean_ws(values.get("ho_ten")),
            "cccd": clean_ws(values.get("dinh_danh_ca_nhan")),
            "values": values,
        })

    return {
        "filename": filename,
        "unit_label": unit_label,
        "file_bytes": file_bytes,
        "layout_note": layout_note,
        "code_row": code_row,
        "data_start_row": data_start_row,
        "col_map": col_map,
        "rows": rows,
    }


def merge_mau03_files(sources):
    """sources: list các dict trả về từ read_source_file(), THEO ĐÚNG THỨ TỰ đã tải lên.
    File ĐẦU TIÊN trong danh sách được dùng làm NỀN — giữ nguyên định dạng, các sheet danh mục
    (NgheNghiep, NoiLamViec, DoiTuongKham...) và 4 dòng đầu; các sheet danh mục của các file khác
    KHÔNG được dùng tới, vì đó là danh mục dùng chung của Medinet chứ không phải dữ liệu riêng của
    từng đơn vị.

    Trả về (merged_bytes, summary):
      summary = {
        "total_rows": tổng số bệnh nhân đã ghép,
        "per_file": [(filename, unit_label, số dòng lấy được), ...],
        "duplicates": [{"cccd": ..., "rows": [(dòng Excel sau ghép, đơn vị, họ tên), ...]}, ...],
        "field_gaps": [(filename, (mã field thiếu so với file nền, ...)), ...],
        "layout_notes": [(filename, ghi chú cấu trúc), ...],
      }
    """
    if len(sources) < 2:
        raise ValueError("Cần ít nhất 2 file để ghép.")

    base = sources[0]
    base_col_map = base["col_map"]
    base_data_start_row = base["data_start_row"]

    wb = openpyxl.load_workbook(BytesIO(base["file_bytes"]))  # không data_only — giữ định dạng để lưu lại
    ws = wb[SHEET_MAIN]

    # Tắt khoá bảo vệ (Protect Sheet) mang theo từ file nền — cùng lỗi và cùng cách vá như
    # annotate_workbook() trong validators.py: không tắt thì Excel tự ẩn/xám bớt nút Home/Filter
    # ở file ghép tải về. Chỉ đổi thuộc tính bảo vệ hiển thị, không đụng cấu trúc/dữ liệu nên
    # không ảnh hưởng chuẩn import Medinet.
    for sh in wb.worksheets:
        sh.protection.sheet = False
    try:
        wb.security = None
    except Exception:
        pass

    # Xoá sạch vùng dữ liệu cũ (kể cả của chính file nền) — sẽ ghi lại TOÀN BỘ theo đúng thứ tự
    # ghép bên dưới, tránh sót dòng thừa nếu tổng số dòng ghép ít hơn số dòng gốc của file nền.
    old_last_row = find_last_data_row(
        ws, [c for c in (base_col_map.get("ho_ten"), base_col_map.get("dinh_danh_ca_nhan")) if c],
        base_data_start_row,
    )
    for r in range(base_data_start_row, old_last_row + 1):
        for c in base_col_map.values():
            ws.cell(row=r, column=c).value = None

    field_gaps = []
    all_rows = []  # (unit_label, filename, ho_ten, cccd, values)
    for src in sources:
        gaps = tuple(sorted(set(base_col_map) - set(src["col_map"])))
        if gaps:
            field_gaps.append((src["filename"], gaps))
        for row in src["rows"]:
            all_rows.append((src["unit_label"], src["filename"], row["ho_ten"], row["cccd"], row["values"]))

    # Ghi dữ liệu đã ghép vào sheet chính, đánh lại STT liên tục — đọc/ghi theo MÃ FIELD nên vẫn
    # đúng cột dù vị trí cột ở file nguồn có lệch đôi chút so với file nền.
    cccd_positions = {}
    for idx, (unit_label, filename, ho_ten, cccd, values) in enumerate(all_rows):
        r = base_data_start_row + idx
        ws.cell(row=r, column=1).value = idx + 1  # STT
        for code, col in base_col_map.items():
            ws.cell(row=r, column=col).value = values.get(code)
        if cccd:
            cccd_positions.setdefault(cccd, []).append((r, unit_label, ho_ten))

    # Đánh dấu (KHÔNG chặn) các CCCD trùng — tô vàng + ghi chú, cùng kiểu FILL_WARN app kiểm tra
    # đang dùng, để nhất quán giao diện giữa 2 tính năng.
    cccd_col = base_col_map.get("dinh_danh_ca_nhan")
    duplicates = []
    if cccd_col:
        for cccd, occurrences in cccd_positions.items():
            if len(occurrences) > 1:
                duplicates.append({"cccd": cccd, "rows": occurrences})
                for r, unit_label, ho_ten in occurrences:
                    cell = ws.cell(row=r, column=cccd_col)
                    cell.fill = FILL_WARN
                    others = [f"{u} (dòng {rr})" for rr, u, _ in occurrences if rr != r]
                    cm = Comment(
                        "CCCD trùng với: " + "; ".join(others) + " — vẫn giữ lại cả hai, tự kiểm tra "
                        "xem có phải cùng 1 người bị nhập 2 lần hay trùng ngẫu nhiên trước khi chạy "
                        "script điền Medinet (CCCD trùng thật sẽ bị Medinet từ chối lưu).",
                        "Công cụ ghép file",
                    )
                    cm.width, cm.height = 320, 90
                    cell.comment = cm

    # Thêm 1 sheet phụ ghi rõ nguồn gốc từng dòng — không đụng vào cấu trúc sheet chính, chỉ để
    # tra cứu/đối chiếu sau này (ví dụ biết dòng lỗi nào cần báo lại cho đúng khóm/tổ).
    if "Nguon_Ghep" in wb.sheetnames:
        del wb["Nguon_Ghep"]
    src_ws = wb.create_sheet("Nguon_Ghep")
    src_ws.append(["STT (sau khi ghép)", "Họ tên", "CCCD", "Đơn vị nguồn", "Tên file gốc", "CCCD trùng?"])
    dup_cccds = {d["cccd"] for d in duplicates}
    for idx, (unit_label, filename, ho_ten, cccd, values) in enumerate(all_rows):
        src_ws.append([idx + 1, ho_ten, cccd, unit_label, filename, "Có" if cccd in dup_cccds else ""])

    out = BytesIO()
    wb.save(out)

    summary = {
        "total_rows": len(all_rows),
        "per_file": [(s["filename"], s["unit_label"], len(s["rows"])) for s in sources],
        "duplicates": duplicates,
        "field_gaps": field_gaps,
        "layout_notes": [(s["filename"], s["layout_note"]) for s in sources if s["layout_note"]],
    }
    return out.getvalue(), summary
